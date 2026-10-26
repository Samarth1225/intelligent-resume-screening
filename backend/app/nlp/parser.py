"""
Resume file -> plain text extraction.

Design goal (per spec section 20, Error Handling): a resume that cannot be
parsed, or a PDF with no extractable text (e.g. a scanned image), must never
crash the pipeline. Callers should catch ParsingError and continue, storing
"Not detected" fields for that candidate instead of failing the whole batch.
"""
from pathlib import Path

import pdfplumber
import docx


class ParsingError(Exception):
    """Raised when a resume file cannot be read at all."""


def extract_text_from_pdf(path: str) -> str:
    text_parts = []
    try:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                try:
                    page_text = page.extract_text() or ""
                except Exception:
                    page_text = ""
                text_parts.append(page_text)
    except Exception as exc:
        raise ParsingError(f"Could not open PDF: {exc}") from exc
    return "\n".join(text_parts).strip()


def extract_text_from_docx(path: str) -> str:
    try:
        document = docx.Document(path)
    except Exception as exc:
        raise ParsingError(f"Could not open DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts).strip()


def extract_text(path: str) -> str:
    """Dispatch by file extension. Returns '' (not an exception) if the file
    opens but contains no extractable text (e.g. a scanned/image-only PDF) —
    the caller treats that as "no readable text" and marks fields Not detected.
    """
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    if ext == ".docx":
        return extract_text_from_docx(path)
    raise ParsingError(f"Unsupported file type: {ext}")
